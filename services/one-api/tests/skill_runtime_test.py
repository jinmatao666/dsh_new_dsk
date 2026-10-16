"""技能发布、Windows 回归和 Mac 分支测试；不请求生产服务。"""
import argparse
import ast
import importlib.util
import io
import json
import os
from pathlib import Path
import shutil
import struct
import subprocess
import sys
import tempfile
import threading
import unittest
import wave
import xml.etree.ElementTree as ET
import zipfile
from contextlib import redirect_stdout
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from unittest.mock import patch

sys.dont_write_bytecode = True
SKILLS = Path(__file__).resolve().parents[1] / "web/air/public/skills"
GIS = ["market-gis-geology-analysis", "market-gis-land-use-plan-review", "market-gis-third-survey-analysis"]


def load(skill, name="office_tools.py"):
    path = SKILLS / skill / "scripts" / name
    key = skill.replace("-", "_") + name.replace(".", "_")
    spec = importlib.util.spec_from_file_location(key, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[key] = module
    spec.loader.exec_module(module)
    return module


def result(output):
    lines = [line[14:] for line in output.splitlines() if line.startswith("WANWEI_RESULT=")]
    return json.loads(lines[-1])


class SkillTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="wanwei-skills-")
        self.directory = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def launch(self, skill, arguments, windows=False):
        script = SKILLS / skill / "scripts"
        command = ([shutil.which("powershell"), "-NoProfile", "-ExecutionPolicy", "Bypass", "-File",
                    str(script / "invoke.ps1")] if windows else [sys.executable, str(script / "invoke.py")])
        completed = subprocess.run(command + arguments, capture_output=True, text=True, encoding="utf-8",
                                   errors="replace", timeout=90, env={**os.environ, "PYTHONIOENCODING": "utf-8",
                                                                   "PYTHONDONTWRITEBYTECODE": "1",
                                                                   "DSH_SKILL_NO_AUTO_OPEN": "1"})
        self.assertEqual(completed.returncode, 0, completed.stderr + completed.stdout)
        return result(completed.stdout)

    def test_all_packages_publish_complete_os_entrypoints(self):
        packages = list(SKILLS.glob("*/manifest.json"))
        self.assertEqual(len(packages), 13)
        for manifest in packages:
            with self.subTest(skill=manifest.parent.name):
                data = json.loads(manifest.read_text(encoding="utf-8-sig"))
                self.assertIn("scripts/invoke.py", data["files"])
                for file in data["files"]:
                    self.assertTrue((manifest.parent / file).is_file(), file)
                instructions = (manifest.parent / "SKILL.md").read_text(encoding="utf-8-sig")
                self.assertIn("macOS", instructions)
                self.assertIn("默认以中文", instructions)
                compile((manifest.parent / "scripts/invoke.py").read_text(encoding="utf-8"), "invoke.py", "exec")
                for script in (manifest.parent / "scripts").glob("*.py"):
                    ast.parse(script.read_text(encoding="utf-8-sig"), filename=str(script), feature_version=(3, 9))

    def test_office_entrypoints_do_not_reject_mac_python39_by_version_alone(self):
        versions = {"python-docx": "1.1.0", "pypdf": "6.0.0", "openpyxl": "3.1.0",
                    "Pillow": "12.0.0", "PyMuPDF": "1.24.0"}
        for directory in sorted(SKILLS.glob("office-*")):
            with self.subTest(skill=directory.name):
                entry = load(directory.name, "invoke.py")
                with patch.object(entry.sys, "version_info", (3, 9, 6)), \
                     patch.object(entry.importlib.util, "find_spec", return_value=object()), \
                     patch.object(entry.importlib.metadata, "version", side_effect=versions.get), \
                     patch.object(entry.shutil, "which", return_value="/usr/bin/soffice"):
                    info = entry.environment()
                self.assertEqual(info["skill"], directory.name)

    def test_office_processors_do_not_reject_mac_python39_by_version_alone(self):
        for directory in sorted(SKILLS.glob("office-*")):
            if not (directory / "scripts/office_tools.py").is_file():
                continue
            with self.subTest(skill=directory.name):
                runtime = load(directory.name)
                arguments = argparse.Namespace(command="test", handler=lambda _: {"success": True})
                parser = argparse.Namespace(parse_args=lambda _: arguments)
                output = io.StringIO()
                with patch.object(runtime.sys, "version_info", (3, 9, 6)), \
                     patch.object(runtime, "build_parser", return_value=parser), \
                     patch.object(runtime, "publication_result", side_effect=lambda value: value), \
                     redirect_stdout(output):
                    code = runtime.main([])
                self.assertEqual(code, 0)
                self.assertTrue(result(output.getvalue())["success"])


    def test_office_processing_and_legacy_windows_entrypoints(self):
        from PIL import Image
        from openpyxl import Workbook, load_workbook
        from pypdf import PdfWriter
        image = self.directory / "输入图.png"
        Image.new("RGB", (80, 60), "blue").save(image)
        pdf = self.directory / "输入.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        with pdf.open("wb") as stream:
            writer.write(stream)
        excel = self.directory / "数据.xlsx"
        workbook = Workbook()
        workbook.active.append(["姓名", "数值"])
        workbook.active.append(["小王", 1])
        workbook.save(excel)
        old, new = self.directory / "原文.txt", self.directory / "新文.txt"
        old.write_text("项目金额：100\n原有事项", encoding="utf-8")
        new.write_text("项目金额：200\n新增事项", encoding="utf-8")
        markdown = self.directory / "摘要.md"
        markdown.write_text("# 项目摘要\n\n## 关键要点\n\n- 已明确项目金额。", encoding="utf-8")
        cases = [
            ("office-pdf-organizer", ["merge", "--inputs", str(pdf), str(pdf)], {"pdf"}),
            ("office-pdf-organizer", ["split", "--input", str(pdf), "--ranges", "1"], {"pdf"}),
            ("office-pdf-to-images", ["--input", str(pdf)], {"image"}),
            ("office-images-to-pdf", ["--inputs", str(image)], {"pdf"}),
            ("office-image-optimizer", ["--inputs", str(image), "--format", "webp"], {"image"}),
            ("office-excel-data-cleaner", ["--input", str(excel)], {"xlsx"}),
            ("office-document-compare", ["--original", str(old), "--revised", str(new)], {"docx", "html"}),
            ("office-document-summary", ["render", "--input", str(markdown)], {"docx"}),
            ("office-document-summary", ["extract", "--inputs", str(old)], set()),
            ("office-meeting-minutes", ["render", "--input", str(markdown)], {"docx"}),
        ]
        for windows in ([False, True] if os.name == "nt" else [False]):
            for index, (skill, args, kinds) in enumerate(cases):
                with self.subTest(skill=skill, windows=windows, index=index):
                    output = self.directory / f"结果-{windows}-{index}"
                    data = self.launch(skill, args + ["--output-directory", str(output)], windows)
                    self.assertEqual({item["kind"] for item in data["artifacts"]}, kinds)
                    for artifact in data["artifacts"]:
                        self.assertTrue(Path(artifact["path"]).is_file())
                    if skill == "office-excel-data-cleaner":
                        self.assertIn("处理报告", load_workbook(data["artifacts"][0]["path"]).sheetnames)

    def test_rename_preview_and_chinese_mapping(self):
        source = self.directory / "原名.txt"
        source.write_text("原始内容", encoding="utf-8")
        args = ["--inputs", str(source), "--prefix", "项目", "--output-directory", str(self.directory)]
        preview = self.launch("office-batch-rename", args)
        self.assertEqual(preview["artifacts"], [])
        self.assertTrue(source.is_file())
        applied = self.launch("office-batch-rename", args + ["--apply"])
        self.assertEqual({item["kind"] for item in applied["artifacts"]}, {"file", "csv"})
        mapping = next(item for item in applied["artifacts"] if item["kind"] == "csv")
        self.assertIn("原文件名,新文件名,原路径,新路径", Path(mapping["path"]).read_text(encoding="utf-8-sig"))

    def test_mac_audio_fallback_and_real_duration(self):
        runtime = load("office-meeting-minutes")
        source, target = self.directory / "录音.m4a", self.directory / "标准.wav"
        source.touch()
        calls = []
        def convert(command, **kwargs):
            calls.append(command)
            with wave.open(str(target), "wb") as audio:
                audio.setparams((1, 2, 16000, 0, "NONE", "not compressed"))
                audio.writeframes(b"\0\0" * 32000)
            return subprocess.CompletedProcess(command, 0, "", "")
        with patch.object(runtime.sys, "platform", "darwin"), patch.dict(os.environ, {"WANWEI_FFMPEG_PATH": ""}), \
             patch.object(runtime.shutil, "which", side_effect=lambda name: "/usr/bin/afconvert" if name == "afconvert" else None), \
             patch.object(runtime.subprocess, "run", side_effect=convert):
            self.assertEqual(runtime._normalize_audio(source, target), target)
        self.assertEqual(calls[0][0], "/usr/bin/afconvert")
        args = argparse.Namespace(base_url="", transcription_url="", transcription_model="", transcription_api_key="",
                                  api_key="", audio_segment_seconds=120, transcription_timeout=10, overwrite=False)
        with patch.object(runtime, "_transcribe_audio_chunk", return_value="这是测试转写"):
            transcript = runtime.transcribe_audio(target, self.directory, args)
        self.assertTrue(transcript.is_file())
        self.assertEqual(args._audio_metadata["audioDurationSeconds"], 2)
        self.assertEqual(args._audio_metadata["segmentCount"], 1)

    def test_missing_mac_environment_fails_without_workaround(self):
        entry = load(GIS[0], "invoke.py")
        with patch.object(entry.sys, "platform", "darwin"), \
             patch.object(entry.shutil, "which", return_value=None):
            info = entry.environment()
        self.assertNotIn("powershell", info)
        runtime = load("office-meeting-minutes")
        source = self.directory / "录音.m4a"
        source.touch()
        with patch.object(runtime.sys, "platform", "darwin"), patch.dict(os.environ, {"WANWEI_FFMPEG_PATH": ""}), \
             patch.object(runtime.shutil, "which", return_value=None):
            with self.assertRaisesRegex(runtime.UserError, "afconvert"):
                runtime._normalize_audio(source, self.directory / "输出.wav")

    def test_gis_native_entry_accepts_macos_python39(self):
        for skill in GIS:
            with self.subTest(skill=skill):
                for name in ("invoke.py", "gis_native.py"):
                    source = (SKILLS / skill / "scripts" / name).read_text(encoding="utf-8")
                    ast.parse(source, feature_version=(3, 9))
                entry = load(skill, "invoke.py")
                with patch.object(entry.sys, "platform", "darwin"), \
                     patch.object(entry.sys, "version_info", (3, 9, 6)):
                    self.assertEqual(entry.environment()["skill"], skill)
                with patch.object(entry.sys, "platform", "darwin"), \
                     patch.object(entry.sys, "version_info", (3, 8, 10)):
                    with self.assertRaisesRegex(RuntimeError, "Python 3.9"):
                        entry.environment()

    @unittest.skipUnless(sys.platform == "darwin", "需要 macOS 原生 afconvert")
    def test_native_mac_afconvert(self):
        runtime = load("office-meeting-minutes")
        source, target = self.directory / "原始录音.wav", self.directory / "标准录音.wav"
        with wave.open(str(source), "wb") as audio:
            audio.setparams((2, 2, 44100, 0, "NONE", "not compressed"))
            audio.writeframes(b"\0\0\0\0" * 44100)
        converter = shutil.which("afconvert")
        self.assertIsNotNone(converter)
        with patch.object(runtime.shutil, "which", side_effect=lambda name: converter if name == "afconvert" else None), \
             patch.dict(os.environ, {"WANWEI_FFMPEG_PATH": ""}):
            runtime._normalize_audio(source, target)
        with wave.open(str(target), "rb") as audio:
            self.assertEqual((audio.getnchannels(), audio.getsampwidth(), audio.getframerate()), (1, 2, 16000))
            self.assertAlmostEqual(audio.getnframes() / audio.getframerate(), 1, places=1)

    def test_meeting_only_word_and_duration_metadata(self):
        runtime = load("office-meeting-minutes")
        source = self.directory / "录音.wav"
        with wave.open(str(source), "wb") as audio:
            audio.setparams((1, 2, 16000, 0, "NONE", "not compressed"))
            audio.writeframes(b"\0\0" * 16000)
        output = io.StringIO()
        with patch.object(runtime, "_transcribe_audio_chunk", return_value="会议决定：先开展测试。"), \
             patch.object(runtime, "generate_meeting_minutes", return_value=runtime.meeting_template()), redirect_stdout(output):
            code = runtime.main(["meeting-prepare", "--audio", str(source), "--output-directory", str(self.directory)])
        self.assertEqual(code, 0)
        data = result(output.getvalue())
        self.assertEqual([item["kind"] for item in data["artifacts"]], ["docx"])
        self.assertEqual(data["metadata"]["audioDurationSeconds"], 1)

    def test_mac_word_converter_returns_real_pdf_and_never_overwrites(self):
        entry = load("office-word-to-pdf", "invoke.py")
        source = self.directory / "会议材料.docx"
        source.write_bytes(b"test")
        existing = self.directory / "会议材料.pdf"
        existing.write_bytes(b"original")
        def convert(command, **kwargs):
            from pypdf import PdfWriter
            target = Path(command[command.index("--outdir") + 1]) / (Path(command[-1]).stem + ".pdf")
            writer = PdfWriter()
            writer.add_blank_page(width=595, height=842)
            with target.open("wb") as content:
                writer.write(content)
            return subprocess.CompletedProcess(command, 0, "", "")
        output = io.StringIO()
        with patch.object(entry.subprocess, "run", side_effect=convert), redirect_stdout(output):
            code = entry.word_to_pdf(["--inputs", str(source), "--output-directory", str(self.directory)],
                                     {"libreoffice": "/Applications/LibreOffice.app/Contents/MacOS/soffice"})
        self.assertEqual(code, 0)
        data = result(output.getvalue())
        self.assertEqual(len(data["artifacts"]), 1)
        self.assertTrue(Path(data["artifacts"][0]["path"]).is_file())
        self.assertEqual(existing.read_bytes(), b"original")

    def test_gis_geojson_and_chinese_shape_zip_against_local_service(self):
        requests = []
        class Handler(BaseHTTPRequestHandler):
            def do_POST(self):
                requests.append(json.loads(self.rfile.read(int(self.headers["Content-Length"]))))
                payload = json.dumps({"YZT_DZHJTJ_LIST": [{"DZHJTJ": "复杂区", "DJ": "3", "ZYMJ": 1}],
                                      "YZT_DZZHYFQK_LIST": [{"FQMC": "低易发区", "DJ": "2", "ZYMJ": 1}],
                                      "YZT_DKDLMJB": [{"DLBM": "0101", "DLMC": "水田", "ZYMJ": 1}],
                                      "YZT_TDFLMJB_HZB": [{"HJMJ": 10, "NYDMJ": 8, "GDMJ": 3,
                                                          "JSYDMJ": 2, "WLYDMJ": 0, "JBNTMJ": 1}],
                                      "YZT_GHSCB": [{"YDZMJ": 10, "SFZYJBNT": True,
                                                     "JBNTMJ": 1, "XZJSQMJ": 2}]},
                                     ensure_ascii=False).encode("utf-8")
                self.send_response(200)
                self.send_header("Content-Type", "application/json; charset=utf-8")
                self.send_header("Content-Length", str(len(payload)))
                self.end_headers()
                self.wfile.write(payload)
            def log_message(self, *args):
                pass
        server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        geojson = self.directory / "地块.geojson"
        points = [[0, 0], [1, 0], [1, 1], [0, 1], [0, 0]]
        geojson.write_text(json.dumps({"type": "Polygon", "coordinates": [points]}), encoding="utf-8")
        content = struct.pack("<i4d2i", 5, 0, 0, 1, 1, 1, 5) + struct.pack("<i", 0)
        content += b"".join(struct.pack("<2d", *point) for point in points)
        header = bytearray(100)
        struct.pack_into(">i", header, 0, 9994)
        shape = bytes(header) + struct.pack(">2i", 1, len(content) // 2) + content
        dbf = bytearray(33)
        struct.pack_into("<H", dbf, 8, 33)
        archive = self.directory / "中文地块.zip"
        with zipfile.ZipFile(archive, "w") as output:
            for suffix, data in [(".shp", shape), (".shx", bytes(100)), (".dbf", dbf),
                                 (".prj", b'GEOGCS["WGS 84"]')]:
                output.writestr("地块一" + suffix, data)
        try:
            for skill in GIS:
                native = load(skill, "gis_native.py")
                for source in [geojson, archive]:
                    with self.subTest(skill=skill, native=True, source=source.name):
                        from docx import Document
                        from openpyxl import load_workbook
                        output = self.directory / f"{skill}-native-{source.suffix[1:]}"
                        data = native.execute({"-InputPath": str(source),
                                               "-OutputDirectory": str(output),
                                               "-BaseUrl": f"http://127.0.0.1:{server.server_port}"})
                        self.assertEqual({item["kind"] for item in data["artifacts"]}, {"docx", "xlsx"})
                        outputs = {item["kind"]: Path(item["path"]) for item in data["artifacts"]}
                        document = Document(outputs["docx"])
                        report = "\n".join(p.text for p in document.paragraphs)
                        self.assertIn("项目与数据概况", report)
                        self.assertIn("核心分析结论", report)
                        self.assertIn("专业研判", report)
                        self.assertIn("实施建议与结论", report)
                        self.assertIn("数据来源与使用限制", report)
                        self.assertNotIn('GEOGCS["', report)
                        self.assertEqual(len(document.tables), 1)
                        if skill == "market-gis-third-survey-analysis":
                            self.assertIn("地类明细", report)
                            self.assertIn("永久基本农田", report)
                        self.assertEqual(load_workbook(outputs["xlsx"]).sheetnames,
                                         ["结果汇总", "专业分析", "接口明细"])
                        sheet = load_workbook(outputs["xlsx"])["结果汇总"]
                        headings = [cell.value for cell in sheet[1]]
                        expected_first = {"market-gis-geology-analysis": "分析图层",
                                          "market-gis-third-survey-analysis": "土地利用指标",
                                          "market-gis-land-use-plan-review": "审查指标"}
                        self.assertEqual(headings[0], expected_first[skill])
                        self.assertGreater(sheet.max_row, 1)
                        for path in outputs.values():
                            with zipfile.ZipFile(path) as package:
                                for member in package.namelist():
                                    if member.endswith(".xml"):
                                        ET.fromstring(package.read(member))
                        self.assertTrue(Path(data["metadata"]["analysisViewPath"]).is_file())
                        self.assertEqual(data["metadata"]["featureCount"], 1)
                for windows in ([False, True] if os.name == "nt" else [False]):
                    for source in [geojson, archive]:
                        with self.subTest(skill=skill, windows=windows, source=source.name):
                            output = self.directory / f"{skill}-{windows}-{source.suffix[1:]}"
                            data = self.launch(skill, ["-InputPath", str(source), "-OutputDirectory", str(output),
                                                      "-BaseUrl", f"http://127.0.0.1:{server.server_port}"], windows)
                            self.assertEqual({item["kind"] for item in data["artifacts"]}, {"docx", "xlsx"})
                            self.assertTrue(Path(data["metadata"]["analysisViewPath"]).is_file())
                            self.assertEqual(data["metadata"]["featureCount"], 1)
            self.assertEqual(len(requests), 18 if os.name == "nt" else 12)
        finally:
            server.shutdown()
            server.server_close()
            thread.join()


if __name__ == "__main__":
    unittest.main()
