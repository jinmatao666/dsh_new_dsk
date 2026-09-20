"""End-to-end smoke tests for the deterministic office skill runtime."""

from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import sys
import tempfile
import unittest
import wave
from pathlib import Path
from unittest import mock

import office_tools


RUNTIME = Path(__file__).with_name("office_tools.py")


class OfficeToolsSmokeTest(unittest.TestCase):
    """Exercise each local processing path with real office artifacts."""

    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory(prefix="wanwei-office-skills-")
        self.root = Path(self.temporary.name)
        self.inputs = self.root / "inputs"
        self.outputs = self.root / "outputs"
        self.inputs.mkdir()
        self.outputs.mkdir()
        self._create_fixtures()

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def run_tool(self, *arguments: str, expected_code: int = 0) -> dict:
        environment = os.environ.copy()
        environment["PYTHONIOENCODING"] = "utf-8"
        completed = subprocess.run(
            [sys.executable, str(RUNTIME), *arguments],
            capture_output=True,
            text=True,
            encoding="utf-8",
            env=environment,
            check=False,
        )
        self.assertEqual(expected_code, completed.returncode, completed.stderr)
        if expected_code != 0:
            return {"stderr": completed.stderr}
        marker = next(line for line in completed.stdout.splitlines() if line.startswith("WANWEI_RESULT="))
        result = json.loads(marker.removeprefix("WANWEI_RESULT="))
        for artifact in result.get("artifacts", []):
            self.assertTrue(Path(artifact["path"]).is_file(), artifact)
        return result

    def test_pdf_image_and_document_workflows(self) -> None:
        image_pdf = self.run_tool(
            "images-to-pdf", "--inputs", str(self.inputs / "red.png"), str(self.inputs / "blue.jpg"),
            "--page-size", "A4", "--orientation", "auto", "--output-dir", str(self.outputs / "image-pdf"),
        )
        pdf = Path(image_pdf["artifacts"][0]["path"])
        merged = self.run_tool(
            "pdf-merge", "--inputs", str(pdf), str(pdf), "--output-dir", str(self.outputs / "merge"),
        )
        merged_pdf = Path(merged["artifacts"][0]["path"])
        self.run_tool(
            "pdf-split", "--input", str(merged_pdf), "--ranges", "1-2,4", "--separate-pages",
            "--output-dir", str(self.outputs / "split"),
        )
        converted = self.run_tool(
            "pdf-to-images", "--input", str(merged_pdf), "--pages", "2-3", "--format", "jpg",
            "--dpi", "96", "--output-dir", str(self.outputs / "pdf-images"),
        )
        self.assertEqual(2, len(converted["artifacts"]))
        self.assertTrue(all(artifact["kind"] == "image" for artifact in converted["artifacts"]))

        cleaned = self.run_tool(
            "excel-process", "--input", str(self.inputs / "ledger.xlsx"), "--dedupe-columns", "编号",
            "--filter", "状态=有效", "--output-dir", str(self.outputs / "excel"),
        )
        self.assertEqual(3, cleaned["rows_before"])
        self.assertEqual(1, cleaned["rows_after"])
        self.assertEqual(["xlsx"], [artifact["kind"] for artifact in cleaned["artifacts"]])

        compared = self.run_tool(
            "document-compare", "--original", str(self.inputs / "old.docx"),
            "--revised", str(self.inputs / "new.docx"), "--output-dir", str(self.outputs / "compare"),
        )
        self.assertEqual(1, compared["counts"]["modified"])
        self.assertEqual(1, compared["counts"]["added"])
        self.assertEqual(0, compared["counts"]["deleted"])
        self.assertEqual(["docx", "html"], [artifact["kind"] for artifact in compared["artifacts"]])
        html_text = Path(compared["artifacts"][1]["path"]).read_text(encoding="utf-8")
        self.assertIn(".diff_next{display:none}", html_text)
        extracted = self.run_tool(
            "document-extract", "--inputs", str(self.inputs / "new.docx"), str(self.inputs / "ledger.xlsx"),
            "--output-dir", str(self.outputs / "extract"),
        )
        self.assertEqual([], extracted["artifacts"])
        self.assertEqual(2, len(extracted["workingFiles"]))
        self.assertTrue(all(Path(path).is_file() for path in extracted["workingFiles"]))
        self.run_tool(
            "render-markdown-docx", "--input", str(self.inputs / "summary.md"),
            "--title", "测试摘要", "--output-dir", str(self.outputs / "render"),
        )

    def test_safe_rename_image_optimization_and_meeting(self) -> None:
        rename_a = self.inputs / "rename-a.txt"
        rename_b = self.inputs / "rename-b.txt"
        rename_a.write_text("a", encoding="utf-8")
        rename_b.write_text("b", encoding="utf-8")
        preview = self.run_tool(
            "batch-rename", "--inputs", str(rename_a), str(rename_b), "--project", "演示",
            "--template", "{project}_{index}_{name}{ext}", "--preview", "--output-dir", str(self.inputs),
        )
        self.assertFalse(preview["applied"])
        self.assertTrue(rename_a.exists())
        applied = self.run_tool(
            "batch-rename", "--inputs", str(rename_a), str(rename_b), "--project", "演示",
            "--template", "{project}_{index}_{name}{ext}", "--apply", "--output-dir", str(self.inputs),
        )
        self.assertTrue(applied["applied"])
        self.assertFalse(rename_a.exists())
        self.assertEqual(2, len(applied["artifacts"]))
        self.assertTrue(all(artifact["kind"] == "file" for artifact in applied["artifacts"]))
        self.assertTrue(all(Path(artifact["path"]).is_file() for artifact in applied["artifacts"]))
        self.assertEqual(2, len(applied["internalFiles"]))

        optimized = self.run_tool(
            "image-process", "--inputs", str(self.inputs / "red.png"), "--format", "webp",
            "--quality", "80", "--max-width", "320", "--output-dir", str(self.outputs / "images"),
        )
        self.assertEqual("image", optimized["artifacts"][0]["kind"])
        self.assertEqual(1, len(optimized["artifacts"]))
        self.assertIn("size_reduction_percent", optimized)
        self.assertEqual(1, len(optimized["files"]))

        meeting = self.run_tool(
            "meeting-prepare", "--transcript", str(self.inputs / "transcript.txt"),
            "--materials", str(self.inputs / "new.docx"), "--skip-synthesis",
            "--meeting-title", "项目联调会议",
            "--output-dir", str(self.outputs / "meeting"),
        )
        self.assertEqual(["docx"], [artifact["kind"] for artifact in meeting["artifacts"]])
        self.assertEqual([], list((self.outputs / "meeting").glob("*.md")))
        from docx import Document
        from docx.oxml.ns import qn

        minutes = Document(meeting["artifacts"][0]["path"])
        self.assertAlmostEqual(21.0, minutes.sections[0].page_width.cm, places=1)
        self.assertAlmostEqual(29.7, minutes.sections[0].page_height.cm, places=1)
        self.assertEqual("会议纪要", minutes.paragraphs[0].text)
        self.assertEqual("项目联调会议", minutes.paragraphs[1].text)
        self.assertTrue(any(paragraph.style.name == "List Bullet" for paragraph in minutes.paragraphs))
        self.assertEqual([2, 5], [len(table.columns) for table in minutes.tables])
        self.assertEqual([3.0, 12.0], [round(cell.width.cm, 1) for cell in minutes.tables[0].rows[0].cells])
        header_properties = minutes.tables[1].rows[0]._tr.get_or_add_trPr()
        self.assertIsNotNone(header_properties.find(qn("w:tblHeader")))

    def test_audio_is_split_and_transcribed_in_order(self) -> None:
        audio = self.inputs / "audio.wav"
        arguments = argparse.Namespace(
            base_url="http://example.invalid/v1",
            transcription_url="",
            transcription_model="",
            transcription_api_key="",
            api_key="",
            audio_segment_seconds=30,
            transcription_timeout=60,
            overwrite=False,
        )
        with mock.patch.object(office_tools, "_transcribe_audio_chunk", side_effect=lambda path, *_: path.stem):
            transcript = office_tools.transcribe_audio(audio, self.outputs, arguments)
        text = transcript.read_text(encoding="utf-8")
        self.assertEqual(3, len(re.findall(r"segment-\d{4}", text)))
        self.assertEqual([], list(self.outputs.glob("wanwei-meeting-temp-*")))

    def test_qwen_audio_uses_native_json_request(self) -> None:
        audio = self.inputs / "audio.wav"

        class Response:
            def __enter__(self):
                return self

            def __exit__(self, *_):
                return False

            def read(self) -> bytes:
                return json.dumps({"output": {"text": "会议转写内容"}}).encode("utf-8")

        with mock.patch.object(office_tools.urllib.request, "urlopen", return_value=Response()) as urlopen:
            text = office_tools._transcribe_audio_chunk(
                audio,
                "https://workspace.cn-beijing.maas.aliyuncs.com/api/v1/services/aigc/multimodal-generation/generation",
                "qwen-audio-3.0-asr-flash",
                "test-key",
                60,
            )
        self.assertEqual("会议转写内容", text)
        request = urlopen.call_args.args[0]
        payload = json.loads(request.data.decode("utf-8"))
        self.assertEqual("qwen-audio-3.0-asr-flash", payload["model"])
        self.assertTrue(payload["input"]["messages"][0]["content"][0]["input_audio"]["data"].startswith("data:audio/wav;base64,"))
        self.assertEqual("Bearer test-key", request.get_header("Authorization"))

    def test_qwen_audio_uses_multipart_request_through_one_api(self) -> None:
        audio = self.inputs / "audio.wav"

        class Response:
            def __enter__(self):
                return self

            def __exit__(self, *_):
                return False

            def read(self) -> bytes:
                return json.dumps({"text": "网关转写内容"}).encode("utf-8")

        with mock.patch.object(office_tools.urllib.request, "urlopen", return_value=Response()) as urlopen:
            text = office_tools._transcribe_audio_chunk(
                audio,
                "http://one-api.example/v1/audio/transcriptions",
                "qwen-audio-3.0-asr-flash",
                "test-key",
                60,
            )
        self.assertEqual("网关转写内容", text)
        request = urlopen.call_args.args[0]
        self.assertTrue(request.get_header("Content-type").startswith("multipart/form-data; boundary="))
        self.assertIn(b'qwen-audio-3.0-asr-flash', request.data)
        self.assertIn(b'filename="audio.wav"', request.data)
        self.assertEqual("Bearer test-key", request.get_header("Authorization"))

    def _create_fixtures(self) -> None:
        from PIL import Image
        from docx import Document
        from openpyxl import Workbook

        Image.new("RGB", (640, 480), "#ef4444").save(self.inputs / "red.png")
        Image.new("RGB", (480, 640), "#2563eb").save(self.inputs / "blue.jpg", quality=90)

        old = Document()
        old.add_heading("项目周报", level=1)
        old.add_paragraph("计划于周一完成初稿。")
        old.save(self.inputs / "old.docx")
        new = Document()
        new.add_heading("项目周报", level=1)
        new.add_paragraph("计划于周三完成终稿。")
        new.add_paragraph("风险：接口仍待确认。")
        new.save(self.inputs / "new.docx")

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "台账"
        sheet.append(["编号", "名称", "状态"])
        sheet.append([1, " 甲项目 ", "有效"])
        sheet.append([1, "甲项目", "有效"])
        sheet.append([2, "乙项目", "停用"])
        workbook.save(self.inputs / "ledger.xlsx")

        (self.inputs / "summary.md").write_text(
            "# 执行摘要\n\n## 关键结论\n\n- 已完成测试。\n\n"
            "| 事项 | 责任人 | 截止时间 |\n| --- | --- | --- |\n| 联调 | 张三 | 9月20日 |\n",
            encoding="utf-8",
        )
        (self.inputs / "transcript.txt").write_text(
            "会议决定周三完成终稿。张三负责接口联调，截止 9 月 20 日。",
            encoding="utf-8",
        )
        with wave.open(str(self.inputs / "audio.wav"), "wb") as audio:
            audio.setnchannels(1)
            audio.setsampwidth(2)
            audio.setframerate(16000)
            audio.writeframes(b"\0\0" * 16000 * 65)


if __name__ == "__main__":
    unittest.main()
