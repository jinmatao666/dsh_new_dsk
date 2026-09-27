import base64
from http.cookies import SimpleCookie
import io
import json
import tempfile
import threading
import time
import unittest
from unittest.mock import patch
import urllib.error
import urllib.request
from docx import Document
from server import create_server


class ServerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.server = create_server(self.temp.name, ("127.0.0.1", 0),
                                    redeem_ticket=lambda ticket: {"user_id": "a" if ticket[0] == "a" else "b"},
                                    public_url="https://expert.test")
        self.thread = threading.Thread(target=self.server.serve_forever)
        self.thread.start()
        self.url = f"http://127.0.0.1:{self.server.server_port}"

    def tearDown(self):
        self.server.shutdown()
        self.thread.join()
        self.server.server_close()
        self.server.worker.close()
        self.temp.cleanup()

    def request(self, path, body=None, cookie=None, origin=None):
        headers = {"Content-Type": "application/json"}
        if cookie:
            headers["Cookie"] = cookie
        if origin:
            headers["Origin"] = origin
        request = urllib.request.Request(self.url + path, data=json.dumps(body).encode() if body is not None else None,
                                         headers=headers)
        try:
            with urllib.request.urlopen(request) as response:
                return response.status, response.read(), response.headers
        except urllib.error.HTTPError as response:
            return response.code, response.read(), response.headers

    def login(self, letter):
        status, _, headers = self.request("/api/session", {"ticket": letter * 64})
        self.assertEqual(status, 200)
        cookie = SimpleCookie(headers["Set-Cookie"])
        self.assertTrue(cookie["expert_document_session"]["secure"])
        return "expert_document_session=" + cookie["expert_document_session"].value

    def test_authentication_origin_and_static_headers(self):
        self.assertEqual(self.request("/api/tasks")[0], 401)
        self.assertEqual(self.request("/api/session", {"ticket": "a" * 64}, origin="https://evil.test")[0], 403)
        status, data, headers = self.request("/")
        self.assertEqual(status, 200)
        self.assertIn("文档智能处理".encode(), data)
        self.assertIn("frame-src 'none'", headers["Content-Security-Policy"])

    def test_upload_path_names_rejected(self):
        cookie = self.login("a")
        tool = "compare"
        name = "../old.txt"
        status, _, _ = self.request("/api/tasks", {"tool": tool, "options": {}, "files": [
            {"name": name, "content": base64.b64encode(b"dummy").decode()}]}, cookie)
        self.assertEqual(status, 400)

    def test_real_task_download_and_user_isolation(self):
        cookie_a, cookie_b = self.login("a"), self.login("b")
        status, data, _ = self.request("/api/tasks", {"tool": "compare", "options": {},
                                      "files": [{"name": "old.txt", "content": base64.b64encode("原文".encode()).decode()},
                                                {"name": "new.txt", "content": base64.b64encode("新文".encode()).decode()}]}, cookie_a)
        self.assertEqual(status, 201)
        task_id = json.loads(data)["id"]
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            items = json.loads(self.request("/api/tasks", cookie=cookie_a)[1])["items"]
            if items[0]["state"] not in ("queued", "running"):
                break
            time.sleep(.1)
        self.assertEqual(items[0]["state"], "succeeded")
        self.assertNotIn("user_id", items[0])
        self.assertEqual(json.loads(self.request("/api/tasks", cookie=cookie_b)[1])["items"], [])
        self.assertEqual(self.request(f"/api/tasks/{task_id}/files/0", cookie=cookie_b)[0], 404)
        status, pdf, _ = self.request(f"/api/tasks/{task_id}/files/0", cookie=cookie_a)
        self.assertEqual(status, 200)
        self.assertIn("文档差异", Document(io.BytesIO(pdf)).paragraphs[0].text)
        paragraphs = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(pdf)).paragraphs)
        self.assertIn("原始版本：old.txt", paragraphs)
        self.assertIn("新版本：new.txt", paragraphs)
        self.assertNotIn("000.txt", paragraphs)

    def test_reject_wrong_type_and_invalid_upload(self):
        cookie = self.login("a")
        self.assertEqual(self.request("/api/tasks", {"tool": "compare", "files": [{"name": "a.exe", "content": "AAAA"}]}, cookie)[0], 400)
        self.assertEqual(self.request("/api/tasks", {"tool": "compare", "files": [{"name": "a.pdf", "content": "not-base64"}]}, cookie)[0], 400)

    def test_requested_comparison_analysis_fails_without_model_and_has_no_downloads(self):
        cookie = self.login("a")
        with patch.dict("os.environ", {"EXPERT_MODEL_URL": "", "EXPERT_MODEL": "", "EXPERT_MODEL_KEY": ""}):
            status, data, _ = self.request("/api/tasks", {"tool": "compare", "options": {"scope": "数字变化"},
                "files": [{"name": "old.txt", "content": base64.b64encode("预算 100 万".encode()).decode()},
                          {"name": "new.txt", "content": base64.b64encode("预算 120 万".encode()).decode()}]}, cookie)
            self.assertEqual(status, 201)
            task_id = json.loads(data)["id"]
            deadline = time.monotonic() + 10
            while time.monotonic() < deadline:
                task = next(item for item in json.loads(self.request("/api/tasks", cookie=cookie)[1])["items"]
                            if item["id"] == task_id)
                if task["state"] not in ("queued", "running"):
                    break
                time.sleep(.1)
        self.assertEqual(task["state"], "failed")
        self.assertIn("尚未配置", task["error"])
        self.assertEqual(task["outputs"], [])
        self.assertEqual(self.request(f"/api/tasks/{task_id}/files/0", cookie=cookie)[0], 404)


if __name__ == "__main__":
    unittest.main()
