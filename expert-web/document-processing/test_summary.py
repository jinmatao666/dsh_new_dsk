from pathlib import Path
import json
import tempfile
import unittest
from unittest.mock import patch
from summary import compare_with_focus, summarize_documents, generate
from processors import DocumentError


class SummaryTests(unittest.TestCase):
    def test_word_export_renders_tables_lists_and_source_markers(self):
        from docx import Document
        from processors import write_docx
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / '成果.docx'
            write_docx(['# 综合摘要', '**重要结论** <sup>1</sup>', '- 核对原文',
                        '| 时间 | 事件 |', '| --- | --- |', '| 2026年 | 审查 |',
                        '<source index="1" name="材料.docx">'], target)
            document = Document(target)
            self.assertEqual(len(document.tables), 1)
            self.assertEqual(document.tables[0].cell(1, 1).text, '审查')
            self.assertTrue(document.paragraphs[1].runs[1].bold)
            text = '\n'.join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn('[1]', text)
            self.assertIn('来源 1：材料.docx', text)
            self.assertNotIn('<sup>', text)
            self.assertNotIn('**', text)

    def test_comparison_uses_both_versions_and_requested_focus(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            old, new = root / "旧.txt", root / "新.txt"
            old.write_text("预算 100 万，截止 10 月 1 日", encoding="utf-8")
            new.write_text("预算 120 万，截止 10 月 8 日", encoding="utf-8")
            prompts = []
            def fixture(prompt):
                prompts.append(prompt)
                return "预算增加 20 万；截止时间延后。"
            outputs = compare_with_focus([old, new], root / "output",
                                        {"scope": "数字变化、日期变化", "requirements": "重点核对预算"}, fixture)
            self.assertIn("重点核对预算", prompts[0])
            self.assertIn(old.read_text(encoding="utf-8"), prompts[0])
            self.assertIn(new.read_text(encoding="utf-8"), prompts[0])
            result = json.loads(outputs[3].read_text(encoding="utf-8"))
            self.assertEqual(result["counts"]["modified"], 1)
            self.assertTrue(result["analysis"]["model_generated"])
            self.assertIn("预算增加", outputs[2].read_text(encoding="utf-8"))

    def test_comparison_invalid_parameters_or_model_failure_leave_no_outputs(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            old, new = root / "旧.txt", root / "新.txt"
            old.write_text("旧文字", encoding="utf-8")
            new.write_text("新文字", encoding="utf-8")
            for options in ({"scope": []}, {"scope": "字" * 2001}):
                with self.assertRaises(DocumentError):
                    compare_with_focus([old, new], root / "output", options)
            with self.assertRaisesRegex(DocumentError, "有效正文"):
                compare_with_focus([old, new], root / "output", {"scope": "变化"}, lambda _: "")
            self.assertFalse((root / "output").exists())

    def test_summary_uses_original_names_without_reading_display_paths(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "000.txt"
            source.write_text("真实材料正文", encoding="utf-8")
            prompts = []
            def fixture(prompt):
                prompts.append(prompt)
                return "已核对材料正文。"
            outputs = summarize_documents([source], root / "output", generator=fixture,
                                          input_names=["项目评审材料.txt"])
            self.assertIn("项目评审材料.txt", prompts[0])
            self.assertNotIn("000.txt", prompts[0])
            metadata = json.loads(outputs[2].read_text(encoding="utf-8"))
            self.assertEqual(metadata["sources"][0]["name"], "项目评审材料.txt")
            self.assertIn("项目评审材料.txt", outputs[1].read_text(encoding="utf-8"))

    def test_source_prompt_and_real_documents(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "材料.txt"
            source.write_text("项目截止时间为 10 月 1 日", encoding="utf-8")
            prompts = []
            def fixture(prompt):
                prompts.append(prompt)
                return "## 时间节点\n10 月 1 日，来源：材料.txt。"
            outputs = summarize_documents([source], root / "output", generator=fixture)
            self.assertEqual(len(outputs), 3)
            self.assertIn(source.read_text(encoding="utf-8"), prompts[0])
            self.assertIn("模型生成", outputs[1].read_text(encoding="utf-8"))

    def test_no_model_is_not_fake_success(self):
        with patch.dict("os.environ", {"EXPERT_MODEL_URL": "", "EXPERT_MODEL_KEY": "", "EXPERT_MODEL": ""}):
            with self.assertRaisesRegex(DocumentError, "尚未配置"):
                generate("材料")

    def test_materials_are_not_silently_truncated(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "材料.txt"
            source.write_text("文" * 60001, encoding="utf-8")
            with self.assertRaisesRegex(DocumentError, "不会静默截断"):
                summarize_documents([source], root / "output", generator=lambda _: "not called")
            self.assertFalse((root / "output").exists())


if __name__ == "__main__":
    unittest.main()
