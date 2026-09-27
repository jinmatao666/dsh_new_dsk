"""Independent document text extraction and directional comparison engine."""
import difflib
import hashlib
import html
import json
from pathlib import Path
import zipfile

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


class DocumentError(ValueError):
    pass


MAX_TEXT = 200_000
TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".tsv", ".json", ".yaml", ".yml"}


def check_archive(path):
    """Reject oversized XML packages before third-party readers expand them."""
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 10000 or sum(entry.file_size for entry in entries) > 64 * 1024 * 1024:
            raise DocumentError("文档解压大小超过处理上限")


def extract_text(path):
    path = Path(path)
    extension = path.suffix.lower()
    if extension in TEXT_EXTENSIONS:
        text = path.read_text(encoding="utf-8-sig")
    elif extension == ".docx":
        check_archive(path)
        document = Document(path)
        blocks = []
        # Preserve paragraph/table order rather than moving all tables to the end.
        for block in document.iter_inner_content():
            if hasattr(block, "rows"):
                blocks.extend(" | ".join(cell.text for cell in row.cells) for row in block.rows)
            else:
                blocks.append(block.text)
        text = "\n".join(blocks)
    elif extension == ".pdf":
        reader = PdfReader(path)
        if reader.is_encrypted or len(reader.pages) > 300:
            raise DocumentError("不支持加密 PDF 或超过 300 页的 PDF")
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    elif extension in (".xlsx", ".xlsm"):
        check_archive(path)
        workbook = load_workbook(path, read_only=True, data_only=True)
        blocks, cells = [], 0
        try:
            for sheet in workbook:
                blocks.append(f"工作表：{sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    cells += len(row)
                    if cells > 100000:
                        raise DocumentError("工作表超过 10 万单元格处理上限")
                    if any(value is not None for value in row):
                        blocks.append(" | ".join("" if value is None else str(value) for value in row))
        finally:
            workbook.close()
        text = "\n".join(blocks)
    else:
        raise DocumentError("文件格式不受支持")
    if not text.strip():
        raise DocumentError("未提取到文字；扫描 PDF 请先进行 OCR")
    if len(text) > MAX_TEXT:
        raise DocumentError("材料文字超过 20 万字符，请拆分后提交")
    return text


def compare_text(old, new):
    """First argument is original, second is revised; counts refer to text lines."""
    old_lines, new_lines = old.splitlines(), new.splitlines()
    if len(old_lines) > 10000 or len(new_lines) > 10000:
        raise DocumentError("对比材料超过 1 万行，请拆分后提交")
    matcher = difflib.SequenceMatcher(a=old_lines, b=new_lines, autojunk=False)
    changes, counts = [], dict(added=0, deleted=0, modified=0, unchanged=0)
    for tag, a, b, c, d in matcher.get_opcodes():
        before, after = old_lines[a:b], new_lines[c:d]
        if tag == "equal":
            counts["unchanged"] += b - a
        elif tag == "replace":
            paired = min(len(before), len(after))
            for kind, original, revised in (("modified", before[:paired], after[:paired]),
                                            ("deleted", before[paired:], []), ("added", [], after[paired:])):
                if original or revised:
                    counts[kind] += max(len(original), len(revised))
                    changes.append({"type": kind, "old": original, "new": revised})
        else:
            kind = "added" if tag == "insert" else "deleted"
            counts[kind] += max(len(before), len(after))
            changes.append({"type": kind, "old": before, "new": after})
    return {"counts": counts, "changes": changes}


def write_docx(lines, target):
    document = Document()
    for line in lines:
        if line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        else:
            document.add_paragraph(line)
    document.save(target)


def source_names(inputs, names=None):
    """Resolve display-only names; filesystem operations always use input paths."""
    if names is None:
        return [Path(path).name for path in inputs]
    if not isinstance(names, list) or len(names) != len(inputs) or any(
        not isinstance(name, str) or not name or len(name) > 255 or
        any(ord(char) < 32 for char in name) or "/" in name or "\\" in name
        for name in names
    ):
        raise DocumentError("材料来源名称无效")
    return names


def compare_documents(inputs, output, input_names=None):
    if len(inputs) != 2:
        raise DocumentError("对比必须选择两份文件：原始版本在前，新版本在后")
    inputs = [Path(path) for path in inputs]
    names = source_names(inputs, input_names)
    output = Path(output)
    output.mkdir(parents=True, exist_ok=False)
    old, new = [extract_text(path) for path in inputs]
    result = compare_text(old, new)
    result["sources"] = [{"name": name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()} for path, name in zip(inputs, names)]
    lines = ["# 文档差异对比报告", "", f"原始版本：{names[0]}", f"新版本：{names[1]}", "", "## 差异统计"]
    labels = {"added": "新增", "deleted": "删除", "modified": "修改", "unchanged": "未变化"}
    lines.extend(f"{labels[key]}：{value} 行" for key, value in result["counts"].items())
    for number, change in enumerate(result["changes"], 1):
        lines.extend(["", f"## 差异 {number}：{labels[change['type']]}",
                      "原内容：" + " / ".join(change["old"]), "新内容：" + " / ".join(change["new"])])
    if not result["changes"]:
        lines.append("两份文档提取后的文本内容一致。")
    lines.extend(["", "## 复核说明", "不比较排版、图片、批注或修订痕迹。重要条款、数字、日期和责任分工请回到原文复核。"])
    report = output / "文档差异对比报告.md"
    report.write_text("\n".join(lines), encoding="utf-8")
    metadata = output / "文档差异.json"
    metadata.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    docx = output / "文档差异对比报告.docx"
    write_docx(lines, docx)
    table = difflib.HtmlDiff(wrapcolumn=90).make_table(old.splitlines(), new.splitlines(),
             fromdesc=html.escape(names[0]), todesc=html.escape(names[1]), context=True, numlines=3)
    visual = output / "文档逐行对比.html"
    visual.write_text('<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta http-equiv="Content-Security-Policy" content="default-src \'none\'; style-src \'unsafe-inline\'; base-uri \'none\'"><title>文档逐行对比</title><style>body{font-family:Microsoft YaHei,sans-serif;padding:28px;color:#1f3657}table{border-collapse:collapse;width:100%}td,th{padding:8px;border:1px solid #dce6f3;white-space:pre-wrap;overflow-wrap:anywhere}.diff_add{background:#dff7e8}.diff_sub{background:#ffe5e8}.diff_chg{background:#fff3c7}.diff_next{display:none}</style><h1>文档逐行对比</h1>' + table + '</html>', encoding="utf-8")
    return [docx, visual, report, metadata]
