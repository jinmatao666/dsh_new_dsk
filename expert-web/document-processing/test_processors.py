import json
from pathlib import Path
import tempfile
import unittest
from docx import Document
from openpyxl import Workbook
from processors import extract_text, compare_text, compare_documents, source_names, DocumentError


class DocumentTests(unittest.TestCase):
    def test_display_names_accept_upload_limit_but_never_change_paths(self):
        self.assertEqual(source_names(["000.txt"], ["字" * 251 + ".txt"]), ["字" * 251 + ".txt"])
        for names in ([], ["../other.txt"], ["x\n.txt"], ["字" * 256]):
            with self.assertRaises(DocumentError):
                source_names(["000.txt"], names)

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def test_directional_diff(self):
        result = compare_text("一样\n原始责任人\n删除", "一样\n新的责任人")
        self.assertEqual(result["counts"], dict(added=0, deleted=1, modified=1, unchanged=1))
        self.assertEqual(result["changes"][0]["old"], ["原始责任人"])
        self.assertEqual(compare_text("", "新增")["counts"]["added"], 1)

    def test_docx_preserves_table_order(self):
        path = self.root / "材料.docx"
        document = Document()
        document.add_paragraph("前文")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "表格"
        document.add_paragraph("后文")
        document.save(path)
        self.assertEqual(extract_text(path), "前文\n表格\n后文")

    def test_xlsx_and_empty_text(self):
        workbook = Workbook()
        workbook.active.append(["事项", 123])
        path = self.root / "材料.xlsx"
        workbook.save(path)
        workbook.close()
        self.assertIn("事项 | 123", extract_text(path))
        blank = self.root / "blank.txt"
        blank.write_text(" ", encoding="utf-8")
        with self.assertRaises(DocumentError):
            extract_text(blank)

    def test_real_outputs_and_html_escaping(self):
        old, new = self.root / "旧.txt", self.root / "新.txt"
        old.write_text("原文\n<script>alert(1)</script>", encoding="utf-8")
        new.write_text("新版\n<script>alert(2)</script>", encoding="utf-8")
        outputs = compare_documents([old, new], self.root / "result")
        self.assertEqual(len(outputs), 4)
        self.assertIn("文档差异对比报告", Document(outputs[0]).paragraphs[0].text)
        self.assertNotIn("<script>", outputs[1].read_text(encoding="utf-8"))
        self.assertEqual(json.loads(outputs[3].read_text(encoding="utf-8"))["counts"]["modified"], 2)


if __name__ == "__main__":
    unittest.main()
