"""Meeting-material synthesis and final Word deliverable owned by this expert."""
from pathlib import Path
import re

from docx import Document
from docx.shared import Cm, Pt
from docx.oxml.ns import qn

from extract import DocumentError, extract_text
from audio import transcribe_audio
from model import generate


def render_word(markdown, target, title, input_names=None):
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Cm(2.3)
    section.left_margin = section.right_margin = Cm(2.5)
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.line_spacing = 1.4
    document.add_heading("会议纪要", 0)
    if title:
        document.add_paragraph(title, style="Subtitle")
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if "|" in line and index + 1 < len(lines) and re.match(r"^\|?\s*:?-", lines[index + 1].strip()):
            headers = [cell.strip() for cell in line.strip("|").split("|")]
            table = document.add_table(rows=1, cols=len(headers))
            table.style = "Light Shading Accent 1"
            for cell, text in zip(table.rows[0].cells, headers):
                cell.text = text
            index += 2
            while index < len(lines) and "|" in lines[index]:
                values = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                row = table.add_row()
                for column, cell in enumerate(row.cells):
                    cell.text = values[column] if column < len(values) else ""
                index += 1
            continue
        if line.startswith("# "):
            if line[2:] != "会议纪要":
                document.add_heading(line[2:], 1)
        elif re.match(r"^#{2,3}\s", line):
            document.add_heading(re.sub(r"^#{2,3}\s", "", line), 1)
        elif re.match(r"^[-*]\s", line):
            document.add_paragraph(line[2:], style="List Bullet")
        elif line:
            document.add_paragraph(line.replace("**", ""))
        index += 1
    if input_names:
        document.add_paragraph("材料来源：" + "、".join(input_names))
    document.add_paragraph("复核说明：本纪要由模型依据提供材料生成，重要数字、日期、责任主体和决策请核对原始材料。")
    document.save(target)


def create_minutes(inputs, output, options=None, generator=generate, transcriber=transcribe_audio, input_names=None):
    options = options or {}
    title = options.get("title", "")
    if not isinstance(title, str) or len(title) > 180:
        raise DocumentError("会议名称不得超过 180 字")
    inputs = [Path(path) for path in inputs]
    names = input_names if input_names is not None else [path.name for path in inputs]
    if not isinstance(names, list) or len(names) != len(inputs) or any(
        not isinstance(name, str) or not name or len(name) > 255 or
        "/" in name or "\\" in name or any(ord(char) < 32 for char in name) for name in names
    ):
        raise DocumentError("材料来源名称无效")
    audio = [path for path in inputs if path.suffix.lower() in (".wav", ".mp3", ".m4a")]
    if not inputs or len(audio) > 1:
        raise DocumentError("请提供材料；一次任务最多一个录音文件")
    output = Path(output)
    sections = []
    for path, name in zip(inputs, names):
        text = transcriber(path, output.parent) if path in audio else extract_text(path)
        sections.append(f"材料：{name}\n{text}")
    materials = "\n\n".join(sections)
    if len(materials) > 60000:
        raise DocumentError("会议材料超过 6 万字符，请拆分；不会静默截断")
    markdown = generator(f"会议名称：{title or '未明确'}\n请根据以下材料生成会议纪要。不能补造缺失事实或精确日期。\n\n{materials}")
    if not isinstance(markdown, str) or not markdown.strip():
        raise DocumentError("纪要模型未返回有效正文")
    output.mkdir(parents=True, exist_ok=False)
    target = output / "会议纪要.docx"
    render_word(markdown, target, title, names)
    # Keep the model's actual text for this website's authenticated task view.
    # It is not a published artifact and is never returned in the task list.
    (output.parent / "summary.md").write_text(markdown, encoding="utf-8")
    # Intermediate transcript/material files are not published as deliverables.
    return [target]
