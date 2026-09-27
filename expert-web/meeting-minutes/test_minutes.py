from pathlib import Path
import tempfile
import unittest
import shutil
import wave
from docx import Document
from extract import DocumentError
from minutes import create_minutes
from audio import transcribe_audio


class MinutesTests(unittest.TestCase):
    def test_original_names_are_citations_not_filesystem_paths(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "000.txt"
            source.write_text("责任人张三。", encoding="utf-8")
            def fixture(prompt):
                self.assertIn("项目会转写稿.txt", prompt)
                self.assertNotIn("000.txt", prompt)
                return "## 待办事项\n张三负责核对。"
            outputs = create_minutes([source], root / "output", generator=fixture,
                                     input_names=["项目会转写稿.txt"])
            text = "\n".join(paragraph.text for paragraph in Document(outputs[0]).paragraphs)
            self.assertIn("材料来源：项目会转写稿.txt", text)
            self.assertNotIn("000.txt", text)
            with self.assertRaisesRegex(DocumentError, "来源名称"):
                create_minutes([source], root / "invalid", input_names=["../其他.txt"])

    def test_material_synthesis_and_only_word(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "材料.txt"
            source.write_text("周三开始，责任人张三。", encoding="utf-8")
            def fixture(prompt):
                self.assertIn("周三开始", prompt)
                return "## 会议基本信息\n| 项目 | 内容 |\n| --- | --- |\n| 时间 | 未明确 |\n## 待办事项\n| 序号 | 事项 | 责任人 | 截止时间 | 状态 |\n| --- | --- | --- | --- | --- |\n| 1 | 开始 | 张三 | 周三 | 未开始 |"
            outputs = create_minutes([source], root / "output", {"title": "项目会"}, generator=fixture)
            self.assertEqual([path.name for path in outputs], ["会议纪要.docx"])
            self.assertIn("责任人", (root / "summary.md").read_text(encoding="utf-8"))
            document = Document(outputs[0])
            self.assertEqual(len(document.tables), 2)
            self.assertEqual(document.tables[1].rows[1].cells[3].text, "周三")

    @unittest.skipUnless(shutil.which("ffmpeg"), "requires FFmpeg")
    def test_real_audio_normalization_and_cleanup(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "recording.wav"
            with wave.open(str(source), "wb") as audio:
                audio.setnchannels(2)
                audio.setsampwidth(2)
                audio.setframerate(8000)
                audio.writeframes(b"\0" * 8000 * 4)
            def fixture(path):
                with wave.open(str(path), "rb") as audio:
                    self.assertEqual(audio.getnchannels(), 1)
                    self.assertEqual(audio.getframerate(), 16000)
                return "测试转写替身"
            self.assertEqual(transcribe_audio(source, root, transcriber=fixture), "测试转写替身")
            self.assertEqual(list(root.iterdir()), [source])

    def test_two_recordings_rejected(self):
        with self.assertRaises(DocumentError):
            create_minutes([Path("a.wav"), Path("b.mp3")], Path("not-created"))


if __name__ == "__main__":
    unittest.main()
