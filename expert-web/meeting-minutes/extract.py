"""Independent document text extraction and directional comparison engine."""
from pathlib import Path
import zipfile

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


class DocumentError(ValueError):
    pass


MAX_TEXT = 200_000
TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".tsv", ".json", ".yaml", ".yml"}


def check_archive(path):
    """Reject oversized XML packages before third-party readers expand them."""
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 10000 or sum(entry.file_size for entry in entries) > 64 * 1024 * 1024:
            raise DocumentError("文档解压大小超过处理上限")


def extract_text(path):
    path = Path(path)
    extension = path.suffix.lower()
    if extension in TEXT_EXTENSIONS:
        text = path.read_text(encoding="utf-8-sig")
    elif extension == ".docx":
        check_archive(path)
        document = Document(path)
        blocks = []
        # Preserve paragraph/table order rather than moving all tables to the end.
        for block in document.iter_inner_content():
            if hasattr(block, "rows"):
                blocks.extend(" | ".join(cell.text for cell in row.cells) for row in block.rows)
            else:
                blocks.append(block.text)
        text = "\n".join(blocks)
    elif extension == ".pdf":
        reader = PdfReader(path)
        if reader.is_encrypted or len(reader.pages) > 300:
            raise DocumentError("不支持加密 PDF 或超过 300 页的 PDF")
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    elif extension in (".xlsx", ".xlsm"):
        check_archive(path)
        workbook = load_workbook(path, read_only=True, data_only=True)
        blocks, cells = [], 0
        try:
            for sheet in workbook:
                blocks.append(f"工作表：{sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    cells += len(row)
                    if cells > 100000:
                        raise DocumentError("工作表超过 10 万单元格处理上限")
                    if any(value is not None for value in row):
                        blocks.append(" | ".join("" if value is None else str(value) for value in row))
        finally:
            workbook.close()
        text = "\n".join(blocks)
    else:
        raise DocumentError("文件格式不受支持")
    if not text.strip():
        raise DocumentError("未提取到文字；扫描 PDF 请先进行 OCR")
    if len(text) > MAX_TEXT:
        raise DocumentError("材料文字超过 20 万字符，请拆分后提交")
    return text
