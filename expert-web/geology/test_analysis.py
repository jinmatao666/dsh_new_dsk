import json
from pathlib import Path
import tempfile
import unittest
from docx import Document
from openpyxl import load_workbook
from geometry import SpatialError
from analysis import analyze, request_payload, normalize_response
from test_geometry import RING


class AnalysisTests(unittest.TestCase):
    def test_model_failure_does_not_publish_partial_reports(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "range.json"
            source.write_text(json.dumps({"type": "Polygon", "coordinates": [RING]}), encoding="utf-8")
            from analysis import DATASETS
            def failed(*args):
                raise SpatialError("解读失败")
            with self.assertRaisesRegex(SpatialError, "解读失败"):
                analyze([source], root / "output", requester=lambda _: {key: [] for key in DATASETS},
                        interpreter=failed)
            self.assertFalse((root / "output").exists())

    def test_payload_flags(self):
        payload = request_payload([RING], {})
        self.assertTrue(payload["IsAnaDzhjtj"])
        self.assertFalse(payload["IsAnaXzCoverBp"])
        self.assertEqual(json.loads(payload["GeoJson"])["rings"], [RING])

    def test_real_reports_without_invented_service_data(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "范围.json"
            source.write_text(json.dumps({"type": "Polygon", "coordinates": [RING]}), encoding="utf-8")
            fixture = {"YZT_DZHJTJ_LIST": [{"DZHJTJ": "=dangerous-formula", "DJ": "二级", "ZYMJ": 12.5}], "YZT_DZZHYFQK_LIST": []}
            outputs = analyze([source], root / "output", {"title": "测试项目"}, requester=lambda payload: fixture)
            self.assertEqual(len(outputs), 4)
            self.assertIn("地质条件分析", Document(outputs[0]).paragraphs[0].text)
            workbook = load_workbook(outputs[1])
            try:
                self.assertEqual(workbook.active.cell(2, 3).data_type, "s")
                self.assertEqual(workbook.active.cell(2, 3).value, "=dangerous-formula")
            finally:
                workbook.close()
            self.assertEqual(json.loads(outputs[2].read_text(encoding="utf-8")), fixture)

    def test_wrong_service_response_is_rejected(self):
        with self.assertRaises(SpatialError):
            normalize_response({"error": "bad service"})


if __name__ == "__main__":
    unittest.main()
