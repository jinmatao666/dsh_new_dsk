"""Geology service adapter and genuine office deliverables; no built-in remote address."""
import json
import os
from pathlib import Path
import urllib.request
from urllib.parse import urlsplit

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from geometry import SpatialError, read_dataset


TITLE = "地质条件分析"
DATASETS = {"YZT_DZHJTJ_LIST": "地质环境条件列表", "YZT_DZZHYFQK_LIST": "地质灾害易发区列表"}
FIELDS = {"ID": "记录编号", "DZHJTJ": "地质环境条件", "FQMC": "易发分区名称", "DJ": "分区等级", "ZYMJ": "重叠面积"}


class NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        return None


def service_request(payload):
    endpoint = os.environ.get("EXPERT_GIS_URL", "")
    parsed = urlsplit(endpoint)
    if parsed.scheme not in ("http", "https") or not parsed.netloc or parsed.username or parsed.fragment:
        raise SpatialError("专家服务端尚未配置有效的 GIS 分析地址")
    # HTTP is needed for the existing internal legacy service; the operator must
    # place it behind a trusted network. End users cannot choose this address.
    request = urllib.request.Request(endpoint, data=json.dumps(payload, ensure_ascii=False).encode(),
                                     headers={"Content-Type": "text/plain; charset=utf-8"})
    try:
        with urllib.request.build_opener(NoRedirect).open(request, timeout=120) as response:
            raw = response.read(16 * 1024 * 1024 + 1)
            if len(raw) > 16 * 1024 * 1024:
                raise ValueError("response too large")
        return json.loads(raw.decode("utf-8-sig"))
    except (OSError, ValueError):
        raise SpatialError("GIS 分析服务调用失败或响应无效，未生成替代结果") from None


def request_payload(rings, options):
    field = options.get("zoneField", "分区名称")
    if not isinstance(field, str) or not 1 <= len(field) <= 80:
        raise SpatialError("分区名称字段参数无效")
    return {"GeoJson": json.dumps({"hasZ": False, "hasM": False, "rings": rings}),
            "IsAnaXzCoverBp": False, "IsAnaDzzhyfqk": True, "IsAnaDzhjtj": True, "YfxFieldName": field}


def normalize_response(response):
    if isinstance(response, dict) and "d" in response:
        response = response["d"]
    if isinstance(response, str):
        response = json.loads(response)
    if not isinstance(response, dict) or not any(key in response for key in DATASETS):
        raise SpatialError("分析服务未返回预期的数据集，请检查服务配置和分析范围")
    result = {}
    for key in DATASETS:
        rows = response.get(key, [])
        if isinstance(rows, str):
            rows = json.loads(rows)
        if isinstance(rows, dict):
            rows = [rows]
        if rows is None:
            rows = []
        if not isinstance(rows, list) or len(rows) > 10000 or any(not isinstance(row, dict) for row in rows):
            raise SpatialError("分析数据集格式无效或记录数量超限")
        result[key] = rows
    return result


def analyze(inputs, output, options=None, requester=service_request, interpreter=None):
    options = options or {}
    title = options.get("title", "")
    coordinate = options.get("coordinateSystem", "")
    if not isinstance(title, str) or len(title) > 180 or not isinstance(coordinate, str) or len(coordinate) > 2000:
        raise SpatialError("项目名称或坐标系参数无效")
    rings, source_crs = read_dataset(inputs)
    raw = requester(request_payload(rings, options))
    datasets = normalize_response(raw)
    interpretation = None
    if interpreter is not None:
        interpretation = interpreter(TITLE, title, options, datasets)
        if not isinstance(interpretation, str) or not interpretation.strip():
            raise SpatialError("GIS 解读服务未返回有效正文")
    output = Path(output)
    output.mkdir(parents=True, exist_ok=False)
    original = output / "接口原始结果.json"
    original.write_text(json.dumps(raw, ensure_ascii=False, indent=2), encoding="utf-8")
    workbook = Workbook()
    workbook.remove(workbook.active)
    document = Document()
    document.add_heading(TITLE + "报告", 0)
    if title:
        document.add_paragraph(title)
    document.add_paragraph("坐标说明：" + (coordinate or source_crs) + "；输入坐标未进行转换。")
    document.add_paragraph("分析结果仅依据提交范围与服务返回数据。各图层面积口径独立，不可跨图层相加；面积单位沿用服务原值，须与服务提供方核对。空数据集不自动解释为没有风险。")
    unknown = 0
    view = {"title": TITLE, "project": title, "datasets": []}
    if interpretation is not None:
        view["interpretation"] = {"text": interpretation, "model_generated": True}
        document.add_heading("综合解读（模型生成）", 1)
        for line in interpretation.splitlines():
            document.add_paragraph(line)
        document.add_paragraph("以下解读由模型依据本次数据生成，不改变接口统计；关键判断请由专业人员核对。")
    for key, rows in datasets.items():
        label = DATASETS[key]
        sheet = workbook.create_sheet(label[:31])
        sheet.append(["序号", "项目", "返回值"])
        document.add_heading(label, 1)
        table = document.add_table(rows=1, cols=3)
        table.style = "Light Shading Accent 1"
        for cell, value in zip(table.rows[0].cells, ("序号", "项目", "返回值")):
            cell.text = value
        display = []
        for number, row in enumerate(rows, 1):
            fields = []
            for field, value in row.items():
                if field not in FIELDS:
                    unknown += 1
                    continue
                text = json.dumps(value, ensure_ascii=False) if isinstance(value, (list, dict)) else str(value if value is not None else "未返回")
                fields.append({"label": FIELDS[field], "value": text})
                # Excel formula injection is blocked by explicit string cell type.
                sheet.append([number, FIELDS[field], text])
                sheet.cell(sheet.max_row, 3).data_type = "s"
                cells = table.add_row().cells
                for cell, item in zip(cells, (str(number), FIELDS[field], text)):
                    cell.text = item
            display.append(fields)
        if not rows:
            document.add_paragraph("当前数据集未返回记录，不能据此推断不存在风险。")
        view["datasets"].append({"name": label, "records": display})
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="3476EB")
        sheet.column_dimensions["A"].width = 10
        sheet.column_dimensions["B"].width = 28
        sheet.column_dimensions["C"].width = 55
    if unknown:
        document.add_paragraph(f"服务响应包含 {unknown} 个尚未映射的字段值，未用于中文报告结论；原始响应已完整保存，请专业人员核对。")
    excel = output / (TITLE + "明细.xlsx")
    word = output / (TITLE + "报告.docx")
    workbook.save(excel)
    workbook.close()
    document.save(word)
    presentation = output / "分析结果.json"
    presentation.write_text(json.dumps(view, ensure_ascii=False, indent=2), encoding="utf-8")
    return [word, excel, original, presentation]
