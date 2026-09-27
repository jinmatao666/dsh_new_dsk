import base64
from http.cookies import SimpleCookie
import io
import json
import tempfile
import threading
import time
import unittest
import urllib.error
import urllib.request
from docx import Document
from server import create_server


def spawn_model_fixture(directory, env):
    """Run the real worker/GIS request with an explicit non-production model substitute."""
    import subprocess
    import sys
    from pathlib import Path
    module = str(Path(__file__).resolve().parent)
    worker = str(Path(module) / "worker.py")
    script = ("import sys,runpy; sys.path.insert(0," + repr(module) + "); "
              "import interpretation; interpretation.interpret=lambda *args: '协议测试解读正文'; "
              "runpy.run_path(" + repr(worker) + ",run_name='__main__')")
    return subprocess.Popen([sys.executable, "-c", script, str(directory)],
                            stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
                            text=True, encoding="utf-8", env=env, cwd=directory,
                            start_new_session=__import__("os").name == "posix")


class ServerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.server = create_server(self.temp.name, ("127.0.0.1", 0),
                                    redeem_ticket=lambda ticket: {"user_id": "a" if ticket[0] == "a" else "b"},
                                    public_url="https://expert.test")
        self.thread = threading.Thread(target=self.server.serve_forever)
        self.thread.start()
        self.url = f"http://127.0.0.1:{self.server.server_port}"

    def tearDown(self):
        self.server.shutdown()
        self.thread.join()
        self.server.server_close()
        self.server.worker.close()
        self.temp.cleanup()

    def request(self, path, body=None, cookie=None, origin=None):
        headers = {"Content-Type": "application/json"}
        if cookie:
            headers["Cookie"] = cookie
        if origin:
            headers["Origin"] = origin
        request = urllib.request.Request(self.url + path, data=json.dumps(body).encode() if body is not None else None,
                                         headers=headers)
        try:
            with urllib.request.urlopen(request) as response:
                return response.status, response.read(), response.headers
        except urllib.error.HTTPError as response:
            return response.code, response.read(), response.headers

    def login(self, letter):
        status, _, headers = self.request("/api/session", {"ticket": letter * 64})
        self.assertEqual(status, 200)
        cookie = SimpleCookie(headers["Set-Cookie"])
        self.assertTrue(cookie["expert_planning_review_session"]["secure"])
        return "expert_planning_review_session=" + cookie["expert_planning_review_session"].value

    def test_authentication_origin_and_static_headers(self):
        self.assertEqual(self.request("/api/tasks")[0], 401)
        self.assertEqual(self.request("/api/session", {"ticket": "a" * 64}, origin="https://evil.test")[0], 403)
        status, data, headers = self.request("/")
        self.assertEqual(status, 200)
        self.assertIn("土地利用规划审查".encode(), data)
        self.assertIn("frame-src 'none'", headers["Content-Security-Policy"])

    def test_mismatched_shape_companions_rejected(self):
        cookie = self.login("a")
        def payload(names):
            return {"tool": "analysis", "options": {}, "files": [
                {"name": name, "content": base64.b64encode(b"dummy").decode()} for name in names]}
        for names in (["a.shp", "b.shx", "a.dbf"], ["a.shp", "a.shx", "a.dbf", "a.SHP"],
                      ["a.geojson", "a.shp"], ["../a.shp", "a.shx", "a.dbf"]):
            self.assertEqual(self.request("/api/tasks", payload(names), cookie)[0], 400)
        self.assertEqual(json.loads(self.request("/api/tasks", cookie=cookie)[1])["items"], [])

    def test_unconfigured_gis_fails_and_user_isolation(self):
        cookie_a, cookie_b = self.login("a"), self.login("b")
        polygon = {"type": "Polygon", "coordinates": [[[0, 0], [0, 10], [10, 10], [10, 0], [0, 0]]]}
        status, data, _ = self.request("/api/tasks", {"tool": "analysis", "options": {},
                                      "files": [{"name": "range.geojson", "content": base64.b64encode(json.dumps(polygon).encode()).decode()}]}, cookie_a)
        self.assertEqual(status, 201)
        task_id = json.loads(data)["id"]
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            items = json.loads(self.request("/api/tasks", cookie=cookie_a)[1])["items"]
            if items[0]["state"] not in ("queued", "running"):
                break
            time.sleep(.1)
        self.assertEqual(items[0]["state"], "failed")
        self.assertIn("尚未配置", items[0]["error"])
        self.assertNotIn("user_id", items[0])
        self.assertEqual(json.loads(self.request("/api/tasks", cookie=cookie_b)[1])["items"], [])
        self.assertEqual(self.request(f"/api/tasks/{task_id}/files/0", cookie=cookie_b)[0], 404)
        self.assertEqual(self.request(f"/api/tasks/{task_id}/files/0", cookie=cookie_a)[0], 404)

    def test_http_gis_task_produces_downloadable_word(self):
        from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
        from unittest.mock import patch
        import os
        fixture = {"YZT_GHSCB":[{"YDZMJ":20,"JBNTMJ":0}],"YZT_GNQMJB_LIST":[]}
        received = []
        class GISHandler(BaseHTTPRequestHandler):
            def log_message(self, *args):
                pass
            def do_POST(self):
                received.append(json.loads(self.rfile.read(int(self.headers["Content-Length"]))))
                data = json.dumps(fixture).encode()
                self.send_response(200)
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(data)
        gis = ThreadingHTTPServer(("127.0.0.1", 0), GISHandler)
        thread = threading.Thread(target=gis.serve_forever)
        thread.start()
        try:
            cookie = self.login("a")
            polygon = {"type": "Polygon", "coordinates": [[[0,0],[0,10],[10,10],[10,0],[0,0]]]}
            with patch.dict(os.environ, {"EXPERT_GIS_URL": f"http://127.0.0.1:{gis.server_port}/Analysis.svc/test"}), patch.object(self.server.worker, "_spawn", side_effect=spawn_model_fixture):
                status, data, _ = self.request("/api/tasks", {"tool": "analysis", "options": {},
                    "files": [{"name": "range.geojson", "content": base64.b64encode(json.dumps(polygon).encode()).decode()}]}, cookie)
                self.assertEqual(status, 201)
                task_id = json.loads(data)["id"]
                deadline = time.monotonic() + 10
                while time.monotonic() < deadline:
                    items = json.loads(self.request("/api/tasks", cookie=cookie)[1])["items"]
                    if items[0]["state"] not in ("queued", "running"):
                        break
                    time.sleep(.1)
            self.assertEqual(items[0]["state"], "succeeded")
            self.assertEqual(len(received), 1)
            self.assertIn("GeoJson", received[0])
            status, word, _ = self.request(f"/api/tasks/{task_id}/files/0", cookie=cookie)
            self.assertEqual(status, 200)
            self.assertTrue(Document(io.BytesIO(word)).tables)
            self.assertIn("协议测试解读正文", "\n".join(p.text for p in Document(io.BytesIO(word)).paragraphs))
            summary_index = items[0]["outputs"].index("分析结果.json")
            status, summary_data, _ = self.request(f"/api/tasks/{task_id}/files/{summary_index}", cookie=cookie)
            self.assertEqual(status, 200)
            summary = json.loads(summary_data)
            self.assertTrue(summary["datasets"])
            self.assertEqual(summary["interpretation"]["text"], "协议测试解读正文")
            self.assertTrue(summary["interpretation"]["model_generated"])
            self.assertIn("records", summary["datasets"][0])
            other_cookie = self.login("b")
            self.assertEqual(self.request(f"/api/tasks/{task_id}/files/{summary_index}", cookie=other_cookie)[0], 404)
            with patch.dict(os.environ, {"EXPERT_GIS_URL": f"http://127.0.0.1:{gis.server_port}/Analysis.svc/test",
                                         "EXPERT_MODEL_URL": "", "EXPERT_MODEL": "", "EXPERT_MODEL_KEY": ""}):
                status, body, _ = self.request("/api/tasks", {"tool": "analysis", "options": {},
                    "files": [{"name": "range.geojson", "content": base64.b64encode(json.dumps(polygon).encode()).decode()}]}, cookie)
                self.assertEqual(status, 201)
                failed_id = json.loads(body)["id"]
                deadline = time.monotonic() + 10
                while time.monotonic() < deadline:
                    failed = next(item for item in json.loads(self.request("/api/tasks", cookie=cookie)[1])["items"]
                                  if item["id"] == failed_id)
                    if failed["state"] not in ("queued", "running"):
                        break
                    time.sleep(.1)
            self.assertEqual(failed["state"], "failed")
            self.assertIn("尚未配置", failed["error"])
            self.assertEqual(failed["outputs"], [])
            self.assertEqual(self.request(f"/api/tasks/{failed_id}/files/0", cookie=cookie)[0], 404)
        finally:
            gis.shutdown()
            thread.join()
            gis.server_close()

    def test_reject_wrong_type_and_invalid_upload(self):
        cookie = self.login("a")
        self.assertEqual(self.request("/api/tasks", {"tool": "compare", "files": [{"name": "a.exe", "content": "AAAA"}]}, cookie)[0], 400)
        self.assertEqual(self.request("/api/tasks", {"tool": "compare", "files": [{"name": "a.pdf", "content": "not-base64"}]}, cookie)[0], 400)


if __name__ == "__main__":
    unittest.main()
